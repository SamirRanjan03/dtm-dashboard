import streamlit as st
import pandas as pd
import sqlite3
import seaborn as sns
import matplotlib.pyplot as plt
import io
from pptx import Presentation

# ---------------- PAGE CONFIG ----------------

st.set_page_config(page_title="Bharti AXA DTM Intelligence", layout="wide")

st.markdown("""
<style>
.main {background-color:#F5F7FA;}
h1,h2,h3 {color:#003A8F;}
[data-testid="stSidebar"] {background-color:#003A8F;color:white;}
.stButton>button {background-color:#E60028;color:white;}
</style>
""", unsafe_allow_html=True)

st.title("📊 Bharti AXA Training & Sales Intelligence")

# ---------------- DATABASE ----------------

@st.cache_resource
def connect_db():
    return sqlite3.connect("dtm_database.db", check_same_thread=False)

conn = connect_db()

def safe_pct(a,b):
    return (a/b)*100 if b>0 else 0

# ---------------- LOGIN ----------------

ADMIN_USER="admin"
ADMIN_PASS="admin123"

st.sidebar.header("Login")

username=st.sidebar.text_input("User ID")
password=st.sidebar.text_input("Password",type="password")

role=None

if username==ADMIN_USER and password==ADMIN_PASS:
    role="admin"

try:
    master_login=pd.read_sql("SELECT * FROM dtm_master",conn)
except:
    master_login=pd.DataFrame()

if role!="admin" and not master_login.empty:

    match=master_login[
    (master_login["dtm_name"]==username) &
    (master_login["ecode"].astype(str)==password)
    ]

    if not match.empty:
        role="dtm"

if role is None:
    st.sidebar.error("Invalid Login")
    st.stop()

st.sidebar.success(f"Welcome {username}")

# ---------------- ADMIN DATA UPLOAD ----------------

if role=="admin":

    st.sidebar.header("Upload Base Data")

    master_file=st.sidebar.file_uploader("DTM Master",type=["xlsx"])
    training_file=st.sidebar.file_uploader("Training Data",type=["xlsx"])
    sales_file=st.sidebar.file_uploader("Sales Data",type=["xlsx"])
    hist_file=st.sidebar.file_uploader("Historical Data",type=["xlsx"])

    if master_file:
        pd.read_excel(master_file).to_sql(
        "dtm_master",conn,if_exists="replace",index=False)

    if training_file:
        pd.read_excel(training_file).to_sql(
        "training_data",conn,if_exists="replace",index=False)

    if sales_file:
        pd.read_excel(sales_file).to_sql(
        "product_activity",conn,if_exists="replace",index=False)

    if hist_file:
        pd.read_excel(hist_file).to_sql(
        "historical_data",conn,if_exists="replace",index=False)

# ---------------- USER BULK UPLOAD ----------------

if role=="dtm":

    st.sidebar.header("Upload Activity")

    activity_file=st.sidebar.file_uploader(
    "Upload Excel / CSV",type=["xlsx","csv"])

    if activity_file:

        if activity_file.name.endswith(".csv"):
            df=pd.read_csv(activity_file)
        else:
            df=pd.read_excel(activity_file)

        df["dtm"]=username

        df.to_sql(
        "product_activity",
        conn,
        if_exists="append",
        index=False)

        st.success("Activity Uploaded")

# ---------------- LOAD DATA ----------------

try:
    master=pd.read_sql("SELECT * FROM dtm_master",conn)
    training=pd.read_sql("SELECT * FROM training_data",conn)
    sales=pd.read_sql("SELECT * FROM product_activity",conn)
except:
    st.warning("Admin must upload data")
    st.stop()

data=training.merge(
master,
left_on="dtm",
right_on="dtm_name",
how="left"
)

if role=="dtm":
    data=data[data["dtm_name"]==username]
    sales=sales[sales["dtm"]==username]

# ---------------- KPI ----------------

status=data["irda_status"].value_counts()

appeared=status.get("Appeared",0)
passed=status.get("Passed",0)

pass_pct=safe_pct(passed,appeared)

fls_clear=safe_pct(
data["fls_achieved"].sum(),
data["fls_target"].sum())

activation=safe_pct(
data["la_activated"].sum(),
data["la_target"].sum())

velocity=safe_pct(
data["la_activated"].sum(),
passed)

c1,c2,c3,c4=st.columns(4)

c1.metric("IRDA Pass %",f"{pass_pct:.1f}%")
c2.metric("FLS Clearance %",f"{fls_clear:.1f}%")
c3.metric("LA Activation %",f"{activation:.1f}%")
c4.metric("Activation Velocity",f"{velocity:.1f}%")

# ---------------- TABS ----------------

tab1,tab2,tab3,tab4,tab5,tab6,tab7=st.tabs([
"IRDA Pipeline",
"FLS Probation",
"LA Activation",
"Swabhiman Product",
"Branch Heatmap",
"Historical Analytics",
"Performance Ranking"
])

# ---------------- IRDA ----------------

with tab1:

    funnel=pd.DataFrame({
    "Stage":["Appeared","Passed"],
    "Count":[appeared,passed]
    })

    st.bar_chart(funnel.set_index("Stage"))

# ---------------- FLS ----------------

with tab2:

    data["Shortfall"]=data["fls_target"]-data["fls_achieved"]

    st.dataframe(data[
    ["dtm","branch","fls_target",
    "fls_achieved","Shortfall"]])

# ---------------- LA ----------------

with tab3:

    la=data.groupby("dtm")["la_activated"].sum()

    st.bar_chart(la)

# ---------------- SWABHIMAN ----------------

with tab4:

    swabhiman=sales[sales["product"]=="Swabhiman"]

    sold=swabhiman[swabhiman["outcome"]=="Sold"]

    conversion=safe_pct(len(sold),len(swabhiman))

    m1,m2,m3=st.columns(3)

    m1.metric("Calls",len(swabhiman))
    m2.metric("Premium",swabhiman["premium"].sum())
    m3.metric("Conversion",f"{conversion:.1f}%")

# ---------------- HEATMAP ----------------

with tab5:

    branch=data.groupby("branch").agg({
    "la_activated":"sum",
    "la_target":"sum"
    })

    branch["Activation %"]=safe_pct(
    branch["la_activated"],
    branch["la_target"])

    fig,ax=plt.subplots()

    sns.heatmap(
    branch[["Activation %"]],
    annot=True,
    cmap="RdYlGn",
    ax=ax)

    st.pyplot(fig)

# ---------------- HISTORICAL ----------------

with tab6:

    try:
        hist=pd.read_sql("SELECT * FROM historical_data",conn)
    except:
        hist=None

    if hist is not None:

        hist["DTM_Score"]=(
        hist["IRDA_Pass"]*0.2+
        hist["FLS_Probation_Clear"]*0.2+
        hist["LA_Activated"]*0.25+
        hist["FTM_KPI"]*0.2+
        hist["JFW_Count"]*0.15)

        ranking=hist.groupby("DTM")["DTM_Score"].mean()

        st.bar_chart(ranking)

# ---------------- PERFORMANCE RANKING ----------------

with tab7:

    st.subheader("Top 5 DTMs")

    dtm_perf=data.groupby("dtm").agg({
    "la_activated":"sum",
    "la_target":"sum"
    })

    dtm_perf["Activation %"]=(
    dtm_perf["la_activated"]/
    dtm_perf["la_target"])*100

    top_dtm=dtm_perf.sort_values(
    "Activation %",
    ascending=False).head(5)

    st.dataframe(top_dtm)

    st.subheader("Bottom 5 DTMs")

    bottom_dtm=dtm_perf.sort_values(
    "Activation %",
    ascending=True).head(5)

    st.dataframe(bottom_dtm)

    st.subheader("Top 5 Branches (LA Activation)")

    branch_perf=data.groupby("branch").agg({
    "la_activated":"sum",
    "la_target":"sum"
    })

    branch_perf["Activation %"]=(
    branch_perf["la_activated"]/
    branch_perf["la_target"])*100

    st.dataframe(
    branch_perf.sort_values(
    "Activation %",
    ascending=False).head(5))

    st.subheader("Bottom 5 Branches (LA Activation)")

    st.dataframe(
    branch_perf.sort_values(
    "Activation %",
    ascending=True).head(5))

    st.subheader("Top 5 Branches (Probation)")

    branch_prob=data.groupby("branch").agg({
    "fls_achieved":"sum",
    "fls_target":"sum"
    })

    branch_prob["Probation %"]=(
    branch_prob["fls_achieved"]/
    branch_prob["fls_target"])*100

    st.dataframe(
    branch_prob.sort_values(
    "Probation %",
    ascending=False).head(5))

    st.subheader("Bottom 5 Branches (Probation)")

    st.dataframe(
    branch_prob.sort_values(
    "Probation %",
    ascending=True).head(5))

# ---------------- REPORT ----------------

if st.button("Generate Report"):

    prs=Presentation()

    slide=prs.slides.add_slide(prs.slide_layouts[0])

    slide.shapes.title.text="DTM Performance Report"

    slide.placeholders[1].text="Training Dashboard"

    slide=prs.slides.add_slide(prs.slide_layouts[1])

    slide.shapes.title.text="Key KPIs"

    slide.placeholders[1].text=f"""
IRDA Pass % : {pass_pct:.1f}%

FLS Clearance % : {fls_clear:.1f}%

LA Activation % : {activation:.1f}%
"""

    buf=io.BytesIO()

    prs.save(buf)

    buf.seek(0)

    st.download_button(
    "Download Report",
    buf,
    "DTM_Report.pptx")
